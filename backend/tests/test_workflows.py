import io
import json
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from fastapi.testclient import TestClient

from config import Settings
from main import create_app
from models import PolicyInput
from providers import FX_URL, GROQ_URL, OFAC_URL, parse_sanctions
from store import Store

POLICY = {
    "title": "Test policy — not regulatory guidance",
    "topic": "aml",
    "version": "1.0",
    "content": "All international transaction transfers require verified sender and receiver identity. Sanctions name matches require enhanced due diligence and human review before approval.",
}
TXN = {
    "transaction_id": "TEST-001",
    "amount": "1234.50",
    "currency": "USD",
    "sender": "Example Sender",
    "receiver": "Example Receiver",
    "sender_country": "India",
    "receiver_country": "UK",
    "description": "International transaction for a documented consulting invoice.",
}


def xml_feed(count=1):
    entries = "".join(
        f"<sdnEntry><uid>{i}</uid><lastName>Example Listed Entity {i}</lastName><sdnType>Entity</sdnType><akaList><aka><lastName>Alias Example {i}</lastName></aka></akaList><programList><program>TEST</program></programList></sdnEntry>"
        for i in range(count)
    )
    return f'<sdnList xmlns="urn:ofac"><publshInformation><Publish_Date>10/19/2026</Publish_Date></publshInformation>{entries}</sdnList>'.encode()


@pytest.fixture
def workspace(tmp_path):
    state = {"mode": "valid", "calls": []}

    def handler(request):
        state["calls"].append(str(request.url).split("?")[0])
        if str(request.url).startswith(OFAC_URL):
            if state["mode"] == "ofac_bad_xml":
                return httpx.Response(
                    200,
                    content=b'<!DOCTYPE a [<!ENTITY x "unsafe">]><sdnList>&x;</sdnList>',
                )
            return (
                httpx.Response(503)
                if state["mode"] == "ofac_fail"
                else httpx.Response(200, content=xml_feed(150))
            )
        if str(request.url).startswith(FX_URL):
            if state["mode"] == "fx_invalid_rate":
                return httpx.Response(
                    200,
                    json=[
                        {
                            "date": datetime.now(timezone.utc).date().isoformat(),
                            "base": "INR",
                            "quote": "USD",
                            "rate": "bad",
                        }
                    ],
                )
            if state["mode"] == "fx_fail":
                return httpx.Response(503)
            if state["mode"] == "fx_stale":
                return httpx.Response(
                    200,
                    json=[
                        {
                            "date": "2000-01-01",
                            "base": "INR",
                            "quote": "USD",
                            "rate": 0.012,
                        }
                    ],
                )
            return httpx.Response(
                200,
                json=[
                    {
                        "date": datetime.now(timezone.utc).date().isoformat(),
                        "base": "INR",
                        "quote": "USD",
                        "rate": 0.012,
                    }
                ],
            )
        if str(request.url) == GROQ_URL:
            if state["mode"] == "rate_limit":
                return httpx.Response(429)
            if state["mode"] == "auth_fail":
                return httpx.Response(401)
            if state["mode"] == "network_fail":
                raise httpx.ConnectError("test offline", request=request)
            body = json.loads(request.content)
            context = json.loads(body["messages"][1]["content"])
            chunk = context["policies"][0]
            citations = [
                {
                    "chunk_id": chunk["chunk_id"],
                    "quote": "All international transaction transfers require verified sender and receiver identity.",
                }
            ]
            if state["mode"] != "bad_quote":
                citations[0]["quote"] = next(
                    entry["quote_id"]
                    for entry in context["citation_quotes"]
                    if entry["chunk_id"] == chunk["chunk_id"]
                    and citations[0]["quote"] in entry["text"]
                )
            if state["mode"] == "bad_quote":
                citations[0]["quote"] = (
                    "Fabricated policy quote that is not in the supplied document."
                )
            if state["mode"] == "bad_id":
                citations[0]["chunk_id"] = "not-a-real-policy"
            if "question" in context:
                content = {
                    "answer": "This test policy requires verified identity for both parties to international transfers.",
                    "citations": citations,
                    "limitations": ["Test fixture only."],
                }
            elif body["response_format"]["json_schema"]["name"] == "investigationbrief":
                content = {
                    "summary": "This test investigation requires identity evidence before a human can reach a conclusion.",
                    "recommendation": "review_for_closure"
                    if state["mode"] == "unsafe_clear"
                    else "investigate",
                    "findings": ["Identity records are not supplied."],
                    "missing_evidence": ["Verified sender and receiver identity."],
                    "citations": citations,
                    "limitations": ["Illustrative test policy only."],
                }
            elif body["response_format"]["json_schema"]["name"] == "policyanalysis":
                content = {
                    "summary": "The policy requires documented identity verification for international transfers.",
                    "obligations": [
                        {
                            "obligation": "Verify the identity of sender and receiver before an international transfer.",
                            "control_kind": "kyc",
                            "citation": citations[0],
                            "requires_human_interpretation": True,
                        }
                    ],
                    "limitations": ["Only supplied excerpts were considered."],
                }
            else:
                content = {
                    "verdict": "acceptable"
                    if state["mode"] == "unsafe_clear"
                    else "needs_review",
                    "risk_score": 40,
                    "reasoning": "The supplied test policy requires identity verification; the transaction does not establish completion.",
                    "citations": citations,
                    "risk_factors": ["Identity evidence not provided."],
                    "limitations": ["Test policy, not legal guidance."],
                }
            if state["mode"] == "invalid_score":
                content["risk_score"] = 500
            message = (
                "not JSON" if state["mode"] == "malformed" else json.dumps(content)
            )
            return httpx.Response(
                200,
                json={
                    "choices": [
                        {
                            "finish_reason": "length"
                            if state["mode"] == "truncated"
                            else "stop",
                            "message": {"content": message},
                        }
                    ]
                },
            )
        raise AssertionError(f"Unexpected external call: {request.url}")

    settings = Settings(
        _env_file=None,
        groq_api_key="test-secret-not-real",
        database_path=tmp_path / "test.db",
    )
    app = create_app(settings, httpx.MockTransport(handler), start_worker=False)
    with TestClient(app) as client:
        authenticate(client)
        yield client, app.state.store, state, settings


def prepare(workspace):
    client, store, state, settings = workspace
    assert client.post("/api/policies", json=POLICY).status_code == 201
    assert client.post("/api/sources/ofac/sync").status_code == 200
    return client, store, state, settings


def test_empty_workspace_and_missing_key(tmp_path):
    settings = Settings(
        _env_file=None, groq_api_key="", database_path=tmp_path / "empty.db"
    )
    with TestClient(create_app(settings, start_worker=False)) as client:
        authenticate(client)
        assert client.get("/api/health").json()["ai"]["status"] == "needs_key"
        assert client.get("/api/dashboard").json()["total"] == 0
        assert client.get("/api/policies").json() == []
        response = client.post("/api/transactions/evaluate", json={"transaction": TXN})
        assert response.status_code == 503
        assert client.get("/api/decisions").json() == []


def test_persistent_policy_search_and_update(tmp_path):
    path = tmp_path / "persistent.db"
    first = Store(path)
    item = first.save_policy(PolicyInput(**POLICY))
    second = Store(path)
    assert (
        second.search('"identity" OR * - DROP TABLE policies;')[0]["policy_id"]
        == item["id"]
    )
    assert second.search("identity", "sanctions") == []
    assert second.search("nonexistentsearchword") == []
    updated = PolicyInput(
        **{
            **POLICY,
            "version": "2.0",
            "content": "All international transaction transfers require additional record retention and document preservation.",
        }
    )
    second.save_policy(updated, item["id"])
    assert second.search("identity") == []
    assert second.search("preservation")[0]["version"] == "2.0"


def test_full_feed_and_aliases():
    result = parse_sanctions(xml_feed(175))
    assert result["count"] == 175
    assert result["entries"][-1]["aliases"] == ["Alias Example 174"]
    assert len(result["sha256"]) == 64


def test_success_review_export_and_evidence_survives_update(workspace):
    client, store, state, settings = prepare(workspace)
    response = client.post("/api/transactions/evaluate", json={"transaction": TXN})
    assert response.status_code == 201, response.text
    record = response.json()
    assert record["transaction"]["amount"] == "1234.50"
    assert record["assessment"]["citations"][0]["version"] == "1.0"
    assert FX_URL not in state["calls"]
    client.put(
        "/api/policies/" + store.policies()[0]["id"],
        json={
            **POLICY,
            "version": "2.0",
            "content": "Updated policy for international transactions. Retain all due diligence records securely for further human investigation.",
        },
    )
    review = client.post(
        f"/api/decisions/{record['id']}/reviews",
        json={
            "reviewer": "Test Reviewer",
            "verdict": "flag",
            "notes": "Missing identity verification confirmed during review.",
        },
    )
    assert review.status_code == 201
    report = client.get(f"/api/decisions/{record['id']}/export")
    assert "attachment" in report.headers["content-disposition"]
    exported = report.json()
    assert exported["assessment"]["verdict"] == "needs_review"
    assert exported["assessment"]["citations"][0]["version"] == "1.0"
    assert exported["reviews"][0]["verdict"] == "flag"
    assert client.get("/api/dashboard").json()["pending"] == 0
    assert Store(settings.database_path).decision(record["id"]) is not None
    calls = len(state["calls"])
    assert (
        client.post("/api/transactions/evaluate", json={"transaction": TXN}).status_code
        == 409
    )
    assert len(state["calls"]) == calls


@pytest.mark.parametrize(
    "mode,status,code",
    [
        ("rate_limit", 429, "groq_rate_limited"),
        ("auth_fail", 503, "groq_access_denied"),
        ("network_fail", 502, "invalid_ai_response"),
        ("malformed", 502, "invalid_ai_response"),
        ("truncated", 502, "invalid_ai_response"),
        ("invalid_score", 502, "invalid_ai_response"),
        ("bad_quote", 502, "unverified_citation"),
        ("bad_id", 502, "unverified_citation"),
    ],
)
def test_no_synthetic_decision_on_provider_failure(workspace, mode, status, code):
    client, store, state, _ = prepare(workspace)
    state["mode"] = mode
    response = client.post("/api/transactions/evaluate", json={"transaction": TXN})
    assert response.status_code == status, response.text
    assert response.json()["detail"]["code"] == code
    assert store.decisions() == []
    assert "test-secret-not-real" not in response.text


@pytest.mark.parametrize("mode", ["fx_fail", "fx_stale", "fx_invalid_rate"])
def test_no_currency_substitution(workspace, mode):
    client, store, state, _ = prepare(workspace)
    state["mode"] = mode
    response = client.post(
        "/api/transactions/evaluate", json={"transaction": {**TXN, "currency": "INR"}}
    )
    assert response.status_code == 502
    assert response.json()["detail"]["code"] == "fx_unavailable"
    assert GROQ_URL not in state["calls"]
    assert store.decisions() == []


def test_missing_and_stale_sanctions_stop_evaluation(workspace):
    client, store, state, _ = workspace
    client.post("/api/policies", json=POLICY)
    assert (
        client.post("/api/transactions/evaluate", json={"transaction": TXN}).status_code
        == 409
    )
    client.post("/api/sources/ofac/sync")
    source = store.source("ofac")
    source["fetched_at"] = (datetime.now(timezone.utc) - timedelta(days=2)).isoformat()
    store.save_source("ofac", source)
    assert client.get("/api/sources").json()["ofac"]["status"] == "stale"
    assert (
        client.post("/api/transactions/evaluate", json={"transaction": TXN}).status_code
        == 409
    )
    state["mode"] = "ofac_fail"
    assert client.post("/api/sources/ofac/sync").status_code == 502
    assert client.get("/api/sources").json()["ofac"]["status"] == "stale"
    assert GROQ_URL not in state["calls"]


def test_alias_match_cannot_be_cleared(workspace):
    client, store, state, _ = prepare(workspace)
    state["mode"] = "unsafe_clear"
    response = client.post(
        "/api/transactions/evaluate",
        json={"transaction": {**TXN, "sender": "Alias Example 149"}},
    )
    assert response.status_code == 502
    assert response.json()["detail"]["code"] == "inconsistent_assessment"
    assert store.decisions() == []


def test_query_evidence_and_no_matching_policy(workspace):
    client, _, state, _ = workspace
    assert (
        client.post("/api/query", json={"query": "Identity verification"}).status_code
        == 409
    )
    client.post("/api/policies", json=POLICY)
    result = client.post("/api/query", json={"query": "Identity verification"})
    assert result.status_code == 200, result.text
    assert result.json()["citations"][0]["title"] == POLICY["title"]
    state["mode"] = "bad_quote"
    assert (
        client.post("/api/query", json={"query": "Identity verification"}).status_code
        == 502
    )


@pytest.mark.parametrize(
    "changes",
    [
        {"amount": "-1"},
        {"amount": "1.999"},
        {"currency": "usd"},
        {"description": "short"},
        {"sender": ""},
    ],
)
def test_input_validation(workspace, changes):
    client, _, _, _ = workspace
    assert (
        client.post(
            "/api/transactions/evaluate", json={"transaction": {**TXN, **changes}}
        ).status_code
        == 422
    )


def test_policy_file_formats_and_limits(workspace):
    client, _, _, _ = workspace
    data = {"title": "Uploaded test policy", "topic": "aml", "version": "1.0"}
    assert (
        client.post(
            "/api/policies/file",
            data=data,
            files={"file": ("test.txt", POLICY["content"].encode())},
        ).status_code
        == 201
    )
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = POLICY["content"]
    buffer = io.BytesIO()
    doc.save(buffer)
    assert (
        client.post(
            "/api/policies/file",
            data=data,
            files={"file": ("test.docx", buffer.getvalue())},
        ).status_code
        == 201
    )
    assert (
        client.post(
            "/api/policies/file", data=data, files={"file": ("test.doc", b"not a docx")}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/policies/file", data=data, files={"file": ("test.pdf", b"not a pdf")}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/policies/file",
            data=data,
            files={"file": ("test.txt", b"a" * 5_000_001)},
        ).status_code
        == 413
    )


def test_unknown_api_does_not_serve_frontend(workspace):
    client, _, _, _ = workspace
    assert client.get("/api/missing").status_code == 404
    assert client.get("/api/decisions/missing").status_code == 404
    assert client.get("/api/decisions?offset=-1").status_code == 422


def test_policy_version_conflict(workspace):
    client, store, _, _ = workspace
    response = client.post("/api/policies", json=POLICY)
    policy_id = response.json()["id"]
    assert client.put("/api/policies/" + policy_id, json=POLICY).status_code == 409
    assert store.policy(policy_id)["version"] == "1.0"


def test_live_rate_evidence_preserves_decimal_amount(workspace):
    client, _, _, _ = prepare(workspace)
    response = client.post(
        "/api/transactions/evaluate", json={"transaction": {**TXN, "currency": "INR"}}
    )
    assert response.status_code == 201, response.text
    rate = response.json()["exchange"]
    assert rate["amount_usd"] == "14.81"
    assert rate["rate"] == "0.012"
    assert rate["source"] == FX_URL


def test_unsafe_feed_does_not_replace_snapshot(workspace):
    client, store, state, _ = prepare(workspace)
    checksum = store.source("ofac")["sha256"]
    state["mode"] = "ofac_bad_xml"
    response = client.post("/api/sources/ofac/sync")
    assert response.status_code == 502
    assert store.source("ofac")["sha256"] == checksum


def test_text_pdf_upload(workspace):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(
        ("BT /F1 12 Tf 40 700 Td (" + POLICY["content"] + ") Tj ET").encode("ascii")
    )
    page[NameObject("/Contents")] = content
    buffer = io.BytesIO()
    writer.write(buffer)
    client, _, _, _ = workspace
    response = client.post(
        "/api/policies/file",
        data={"title": "PDF test policy"},
        files={"file": ("policy.pdf", buffer.getvalue())},
    )
    assert response.status_code == 201, response.text
    assert "verified sender" in response.json()["content"]


def authenticate(client):
    hub = client.app.state.hub
    response = client.post(
        "/api/auth/setup",
        json={
            "email": "owner@example.invalid",
            "password": "Test-only-password-2026",
            "name": "Authenticated Owner",
            "workspace": "Test workspace",
            "bootstrap_token": hub.bootstrap_file.read_text(),
        },
    )
    assert response.status_code == 201, response.text
    client.headers["X-CSRF-Token"] = response.json()["csrf"]
    assert (
        client.put(
            "/api/preferences",
            json={
                "enabled_sources": ["ofac"],
                "monitoring_enabled": False,
                "monitoring_interval_hours": 6,
            },
        ).status_code
        == 200
    )
    return response.json()
